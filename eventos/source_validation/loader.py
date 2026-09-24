"""Adaptadores de entrada para inventarios JSON, CSV y Excel."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any

from .models import SourceCandidate


HEADER_MAP = {
    "id": "id",
    "fuente": "name",
    "institución responsable": "institution",
    "institucion responsable": "institution",
    "tipo de fuente": "source_type",
    "cobertura": "coverage",
    "región": "region",
    "region": "region",
    "comuna": "commune",
    "áreas culturales": "categories",
    "areas culturales": "categories",
    "canal de publicación": "channel",
    "canal de publicacion": "channel",
    "url de eventos": "url",
    "url de respaldo / evidencia": "evidence_url",
    "estado de integración": "previous_status",
    "estado de integracion": "previous_status",
    "elegible para porcentaje": "eligible",
    "prioridad": "priority",
    "método potencial": "proposed_method",
    "metodo potencial": "proposed_method",
    "verificada el": "verified_at",
    "observaciones": "notes",
}


def _bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    return str(value or "").strip().casefold() in {"1", "true", "sí", "si", "yes", "on"}


def _candidate(payload: dict[str, Any], index: int) -> SourceCandidate:
    normalized = {
        "id": str(payload.get("id") or f"CL-FUE-{index:04d}"),
        "name": str(payload.get("name") or payload.get("fuente") or "").strip(),
        "institution": str(payload.get("institution") or payload.get("institucion") or "").strip(),
        "source_type": str(payload.get("source_type") or payload.get("tipo_fuente") or "").strip(),
        "coverage": str(payload.get("coverage") or payload.get("cobertura") or "").strip(),
        "region": str(payload.get("region") or "").strip(),
        "commune": str(payload.get("commune") or payload.get("comuna") or "").strip(),
        "categories": str(payload.get("categories") or payload.get("areas") or "").strip(),
        "channel": str(payload.get("channel") or payload.get("canal") or "").strip(),
        "url": str(payload.get("url") or payload.get("url_eventos") or "").strip(),
        "evidence_url": str(payload.get("evidence_url") or payload.get("url_respaldo") or "").strip(),
        "previous_status": str(payload.get("previous_status") or payload.get("estado") or "").strip(),
        "eligible": _bool(payload.get("eligible", payload.get("elegible"))),
        "priority": str(payload.get("priority") or payload.get("prioridad") or "").strip(),
        "proposed_method": str(payload.get("proposed_method") or payload.get("metodo") or "").strip(),
        "verified_at": str(payload.get("verified_at") or payload.get("verificada_el") or "").strip(),
        "notes": str(payload.get("notes") or payload.get("observaciones") or "").strip(),
    }
    if not normalized["name"]:
        raise ValueError(f"La fila {index} no tiene nombre de fuente")
    return SourceCandidate(**normalized)


def load_json(path: Path) -> list[SourceCandidate]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    rows = payload.get("rows") if isinstance(payload, dict) else payload
    if not isinstance(rows, list):
        raise ValueError(f"{path} no contiene una lista ni una clave rows")
    return [_candidate(row, index) for index, row in enumerate(rows, 1) if isinstance(row, dict)]


def load_csv(path: Path) -> list[SourceCandidate]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        rows = []
        for index, raw in enumerate(reader, 1):
            mapped = {HEADER_MAP.get(str(key).strip().casefold(), str(key)): value for key, value in raw.items()}
            rows.append(_candidate(mapped, index))
        return rows


def load_xlsx(path: Path, sheet_name: str = "Fuentes") -> list[SourceCandidate]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Para leer Excel instala openpyxl desde requirements.txt") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[sheet_name]
        iterator = sheet.iter_rows(values_only=True)
        headers = next(iterator)
        mapped_headers = [HEADER_MAP.get(str(value or "").strip().casefold(), "") for value in headers]
        result = []
        for index, values in enumerate(iterator, 1):
            payload = {key: value for key, value in zip(mapped_headers, values) if key}
            if not any(value not in {None, ""} for value in payload.values()):
                continue
            result.append(_candidate(payload, index))
        return result
    finally:
        workbook.close()


def load_candidates(path: Path) -> list[SourceCandidate]:
    suffix = path.suffix.casefold()
    if suffix == ".json":
        candidates = load_json(path)
    elif suffix == ".csv":
        candidates = load_csv(path)
    elif suffix == ".xlsx":
        candidates = load_xlsx(path)
    else:
        raise ValueError(f"Formato no compatible: {path.suffix}")
    identifiers = [candidate.id for candidate in candidates]
    if len(identifiers) != len(set(identifiers)):
        raise ValueError("El inventario contiene IDs duplicados")
    return candidates


def write_canonical_inventory(candidates: list[SourceCandidate], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {"schema_version": 1, "count": len(candidates), "rows": [item.to_dict() for item in candidates]}
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
