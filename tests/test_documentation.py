import importlib.util
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]


def load_tool(name: str):
    path = ROOT / "tools" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    sys.modules[name] = module
    spec.loader.exec_module(module)
    return module


def document_signature(path: Path):
    document = Document(path)
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in document.tables
    ]
    headers = [
        [paragraph.text for paragraph in section.header.paragraphs]
        for section in document.sections
    ]
    footers = [
        [paragraph.text for paragraph in section.footer.paragraphs]
        for section in document.sections
    ]
    return paragraphs, tables, headers, footers


class DatabaseDocumentationTests(unittest.TestCase):
    def test_database_reference_lists_every_versioned_migration(self):
        generator = load_tool("create_database_word")
        text = generator.SOURCE.read_text(encoding="utf-8")
        documented = generator.extract_migration_ids(text)
        actual = [path.name.split("_", 1)[0] for path in sorted(
            (ROOT / "supabase" / "migrations").glob("*.sql")
        )]
        self.assertEqual(documented, actual)

    def test_mermaid_er_is_generated_from_the_complete_dbml(self):
        generator = load_tool("generate_model_er")
        tables = generator.parse_dbml(
            generator.DEFAULT_SOURCE.read_text(encoding="utf-8")
        )
        rendered = generator.render_mermaid(tables)

        self.assertEqual(len(tables), 15)
        self.assertEqual(sum(len(table.columns) for table in tables), 115)
        self.assertEqual(
            generator.DEFAULT_OUTPUT.read_text(encoding="utf-8"),
            rendered,
        )
        for table in tables:
            self.assertIn(f"    {table.name} {{", rendered)
            for column in table.columns:
                self.assertIn(
                    f"        {column.data_type} {column.name}",
                    rendered,
                )

    def test_published_word_is_reproducible_from_database_markdown(self):
        generator = load_tool("create_database_word")
        published = ROOT / "docs" / "Documentacion_Base_de_Datos_Eventos_Culturales.docx"

        with tempfile.TemporaryDirectory() as directory:
            regenerated = Path(directory) / published.name
            previous_output = generator.OUTPUT
            try:
                generator.OUTPUT = regenerated
                generator.build_document()
            finally:
                generator.OUTPUT = previous_output

            self.assertEqual(
                document_signature(published),
                document_signature(regenerated),
            )
            numbering = Document(regenerated).part.numbering_part.element
            saw_instance = False
            for child in numbering:
                if child.tag == qn("w:num"):
                    saw_instance = True
                elif child.tag == qn("w:abstractNum"):
                    self.assertFalse(saw_instance, "Abstract numbering must precede instances")


if __name__ == "__main__":
    unittest.main()
