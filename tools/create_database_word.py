from __future__ import annotations

import re
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "DATABASE.md"
OUTPUT = ROOT / "docs" / "Documentacion_Base_de_Datos_Eventos_Culturales.docx"
DIAGRAM = ROOT / "docs" / "assets" / "database_relationships.png"

NAVY = "17324D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "5E6B78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "C9D3DE"
WHITE = "FFFFFF"
GOLD = "B8871B"
RED = "9B1C1C"

# Preset: compact_reference_guide.
# Named overrides: editorial_cover; technical_table_body (9 pt);
# code_block (Consolas 8.5 pt); relationship_figure.


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths: list[int]) -> None:
    total = sum(widths)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), BORDER)
        borders.append(node)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel_id = part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([props, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str, *, size: float | None = None, color: str | None = None) -> None:
    position = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > position:
            run = paragraph.add_run(text[position : match.start()])
            set_run_font(run, size=size, color=color)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, name="Consolas", size=(size or 10) - 0.3, color=DARK_BLUE)
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "EEF2F6")
            run._element.get_or_add_rPr().append(shading)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, url)
        position = match.end()
    if position < len(text):
        run = paragraph.add_run(text[position:])
        set_run_font(run, size=size, color=color)


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    set_run_font(run, size=8.5, color=MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def add_numbering_definition(document: Document, kind: str) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=-1) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    lvl_text = OxmlElement("w:lvlText")
    lvl_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    justification = OxmlElement("w:lvlJc")
    justification.set(qn("w:val"), "left")
    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "270")
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "300")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.extend([tabs, indent, spacing])
    lvl.extend([start, fmt, lvl_text, justification, p_pr])
    abstract.append(lvl)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def apply_numbering(paragraph, num_id: int) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_element])
    p_pr.append(num_pr)


def create_relationship_diagram(path: Path) -> None:
    width, height = 1600, 880
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    bold_path = Path("C:/Windows/Fonts/arialbd.ttf")
    font = ImageFont.truetype(str(font_path), 30)
    small = ImageFont.truetype(str(font_path), 24)
    bold = ImageFont.truetype(str(bold_path), 34)

    def box(x, y, w, h, title, subtitle, fill=LIGHT_BLUE, outline=BLUE):
        draw.rounded_rectangle((x, y, x + w, y + h), radius=20, fill=f"#{fill}", outline=f"#{outline}", width=4)
        draw.text((x + 22, y + 17), title, font=bold, fill=f"#{NAVY}")
        draw.text((x + 22, y + 64), subtitle, font=small, fill=f"#{MUTED}")

    def arrow(x1, y1, x2, y2, label=""):
        draw.line((x1, y1, x2, y2), fill=f"#{MUTED}", width=5)
        angle = __import__("math").atan2(y2 - y1, x2 - x1)
        length = 18
        for delta in (2.55, -2.55):
            px = x2 + length * __import__("math").cos(angle + delta)
            py = y2 + length * __import__("math").sin(angle + delta)
            draw.line((x2, y2, px, py), fill=f"#{MUTED}", width=5)
        if label:
            tx, ty = (x1 + x2) // 2, (y1 + y2) // 2 - 26
            bbox = draw.textbbox((tx, ty), label, font=small, anchor="mm")
            draw.rectangle((bbox[0] - 6, bbox[1] - 3, bbox[2] + 6, bbox[3] + 3), fill="white")
            draw.text((tx, ty), label, font=small, fill=f"#{MUTED}", anchor="mm")

    box(555, 325, 490, 150, "EVENTS", "Registro canónico consolidado", fill="DDEAF6")
    box(70, 80, 370, 120, "SOURCES", "Procedencia")
    box(615, 80, 370, 120, "ORGANIZERS", "Responsable")
    box(1160, 80, 370, 120, "COMMUNES", "Territorio")
    box(70, 620, 370, 120, "SCRAPE_RUNS", "Bitácora de carga", fill="F4F6F9", outline="7B8794")
    box(615, 620, 370, 120, "MEDIA_ASSETS", "Imágenes y recursos")
    box(1160, 325, 370, 150, "CATEGORIES", "Clasificación N:M")
    box(1160, 650, 370, 150, "APP FUTURA", "Perfiles · favoritos · reportes", fill="FFF4D8", outline=GOLD)

    arrow(255, 200, 555, 350, "1:N")
    arrow(800, 200, 800, 325, "1:N")
    arrow(1345, 200, 1045, 350, "1:N")
    arrow(1045, 400, 1160, 400, "N:M")
    arrow(800, 475, 800, 620, "1:N")
    arrow(440, 650, 590, 470, "sincroniza")
    arrow(1160, 690, 1000, 470, "favoritos / reportes")
    draw.text((800, 835), "Relaciones principales del esquema public", font=font, fill=f"#{NAVY}", anchor="mm")
    image.save(path, quality=95)


def configure_styles(doc: Document) -> None:
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Title": (30, NAVY, 0, 8),
        "Subtitle": (14, MUTED, 0, 10),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def configure_section(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("EVENTOS CULTURALES · DICCIONARIO DE DATOS")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run("Uso personal  ·  Página ")
    set_run_font(run, size=8.5, color=MUTED)
    add_field(p, "PAGE")


def add_cover(doc: Document) -> None:
    for _ in range(4):
        doc.add_paragraph()
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("DOCUMENTACIÓN TÉCNICA")
    set_run_font(run, size=10.5, color=GOLD, bold=True)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Base de datos de\neventos culturales")

    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Diccionario de tablas, columnas, relaciones y categorías")

    line = doc.add_paragraph()
    line.paragraph_format.space_before = Pt(18)
    line.paragraph_format.space_after = Pt(52)
    p_pr = line._p.get_or_add_pPr()
    p_borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "10")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_borders.append(bottom)
    p_pr.append(p_borders)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(5)
    run = meta.add_run("MVP · Agregador de eventos culturales de Chile")
    set_run_font(run, size=11, color=NAVY, bold=True)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(5)
    run = meta.add_run("Versión documentada: 1 de septiembre de 2026")
    set_run_font(run, size=10.5, color=MUTED)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Documento de uso personal")
    set_run_font(run, size=10.5, color=MUTED, italic=True)
    doc.add_page_break()


def add_contents(doc: Document, headings: list[str]) -> None:
    doc.add_heading("Contenido", level=1)
    intro = doc.add_paragraph("Guía de consulta rápida del esquema y su operación.")
    intro.paragraph_format.space_after = Pt(12)
    for index, heading in enumerate(headings, start=1):
        heading = re.sub(r"^\d+\.\s*", "", heading)
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.15)
        p.paragraph_format.space_after = Pt(5)
        run = p.add_run(f"{index:02d}")
        set_run_font(run, size=9.5, color=GOLD, bold=True)
        run = p.add_run(f"   {heading}")
        set_run_font(run, size=11, color=NAVY, bold=True)
    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.12)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.space_before = Pt(12)
    note.paragraph_format.space_after = Pt(0)
    p_pr = note._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), LIGHT_GRAY)
    p_pr.append(shading)
    add_inline(note, "Nota: las categorías son dinámicas y la sección 5 representa una fotografía del conjunto local al 1 de septiembre de 2026.")
    doc.add_page_break()


def parse_markdown_table(lines: list[str]) -> list[list[str]]:
    rows = []
    for line in lines:
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
        rows.append(cells)
    return [rows[0], *rows[2:]]


def choose_widths(headers: list[str]) -> list[int]:
    count = len(headers)
    normalized = [h.lower() for h in headers]
    if count == 4 and "columna" in normalized:
        return [1800, 1400, 1800, 4360]
    if count == 3 and "eventos observados" in normalized:
        return [2500, 1100, 5760]
    if count == 3:
        return [2100, 1500, 5760]
    if count == 2:
        return [2700, 6660]
    return [9360 // count] * count


def add_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    table = doc.add_table(rows=len(rows), cols=len(headers))
    set_table_geometry(table, choose_widths(headers))
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for row_index, values in enumerate(rows):
        for col_index, value in enumerate(values):
            cell = table.cell(row_index, col_index)
            if row_index == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.08
            if len(headers) == 3 and col_index == 1:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline(paragraph, value, size=9, color=NAVY if row_index == 0 else None)
            for run in paragraph.runs:
                if row_index == 0:
                    run.bold = True
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)


def add_code_block(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.15)
    paragraph.paragraph_format.right_indent = Inches(0.15)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.05
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F1F3F5")
    p_pr.append(shading)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(line)
        set_run_font(run, name="Consolas", size=8.5, color=NAVY)


def add_relationship_figure(doc: Document) -> None:
    create_relationship_diagram(DIAGRAM)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(4)
    run = paragraph.add_run()
    inline = run.add_picture(str(DIAGRAM), width=Inches(6.25))
    inline._inline.docPr.set("descr", "Diagrama de relaciones entre fuentes, organizadores, comunas, eventos, categorías, multimedia y tablas futuras.")
    caption = doc.add_paragraph("Figura 1. Relaciones principales del esquema de datos.")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(0)
    caption.paragraph_format.space_after = Pt(10)
    for run in caption.runs:
        set_run_font(run, size=9, color=MUTED, italic=True)


def build_document() -> None:
    text = SOURCE.read_text(encoding="utf-8")
    source_lines = text.splitlines()
    level_two_headings = [line[3:].strip() for line in source_lines if line.startswith("## ")]

    doc = Document()
    configure_styles(doc)
    configure_section(doc)
    doc.core_properties.title = "Base de datos de eventos culturales"
    doc.core_properties.subject = "Diccionario de tablas, columnas, relaciones y categorías"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "Supabase, PostgreSQL, eventos culturales, diccionario de datos"

    bullet_num_id = add_numbering_definition(doc, "bullet")
    decimal_num_id = add_numbering_definition(doc, "decimal")
    add_cover(doc)
    add_contents(doc, level_two_headings)

    lines = source_lines[1:]
    index = 0
    skip_metadata = True
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if skip_metadata and (not stripped or stripped.startswith("**Versión") or stripped.startswith("**Esquema") or stripped.startswith("**Migración")):
            index += 1
            continue
        skip_metadata = False

        if stripped.startswith("```"):
            language = stripped[3:].strip()
            block = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                block.append(lines[index])
                index += 1
            if language == "mermaid":
                add_relationship_figure(doc)
            else:
                add_code_block(doc, block)
            index += 1
            continue

        if stripped.startswith("## "):
            doc.add_heading(stripped[3:], level=1)
        elif stripped.startswith("### "):
            doc.add_heading(stripped[4:], level=2)
        elif stripped.startswith("#### "):
            doc.add_heading(stripped[5:], level=3)
        elif stripped.startswith("| "):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_table(doc, parse_markdown_table(table_lines))
            continue
        elif re.match(r"^- ", stripped):
            paragraph = doc.add_paragraph()
            apply_numbering(paragraph, bullet_num_id)
            add_inline(paragraph, stripped[2:])
        elif re.match(r"^\d+\. ", stripped):
            paragraph = doc.add_paragraph()
            apply_numbering(paragraph, decimal_num_id)
            add_inline(paragraph, re.sub(r"^\d+\. ", "", stripped))
        elif stripped:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, stripped)
        index += 1

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
